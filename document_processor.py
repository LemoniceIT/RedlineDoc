"""
Document text extraction and comparison logic for RedlineDoc.

Extraction strategy (per requirements):
    .doc   -> pywin32 (win32com.client, Word.Application COM)
    .docx  -> python-docx
    .pdf   -> pdfplumber

Comparison:
    difflib.HtmlDiff generates a temporary HTML report highlighting
    differences, additions, and omissions relative to the original document.
"""

from __future__ import annotations

from i18n import t

import difflib
import os
import tempfile
import webbrowser
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ExtractOptions:
    """Options controlling how document text is extracted."""

    ignore_formatting: bool = True
    ignore_images: bool = True
    ignore_header: bool = False
    ignore_footer: bool = False


class DocumentProcessor:
    """Extracts text from supported documents and compares two texts."""

    def __init__(self, log_callback=None) -> None:
        self._log = log_callback or (lambda _msg: None)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    @staticmethod
    def extract_text(file_path: str | Path, options: ExtractOptions) -> str:
        """Extract plain text from a document according to *options*."""
        ext = Path(file_path).suffix.lower()
        if ext == ".docx":
            return DocumentProcessor._extract_docx(file_path, options)
        elif ext == ".doc":
            return DocumentProcessor._extract_doc(file_path, options)
        elif ext == ".pdf":
            return DocumentProcessor._extract_pdf(file_path, options)
        else:
            raise ValueError(t("doc.unsupported_format").format(ext=ext))

    @staticmethod
    def compare_documents(
        original_text: str,
        modified_text: str,
        from_desc: str | None = None,
        to_desc: str | None = None,
        output_path: str | None = None,
    ) -> str:
        """
        Produce an HTML diff report via difflib.HtmlDiff and return the file path.

        The report is written to *output_path* (or a temporary file) and can be
        opened in the default browser by the caller.
        """
        if from_desc is None:
            from_desc = t("doc.original")
        if to_desc is None:
            to_desc = t("doc.modified")

        diff = difflib.HtmlDiff(wrapcolumn=88)

        original_lines = original_text.splitlines()
        modified_lines = modified_text.splitlines()

        html_content = diff.make_file(
            original_lines,
            modified_lines,
            fromdesc=from_desc,
            todesc=to_desc,
            context=True,
        )

        if output_path is None:
            fd, output_path = tempfile.mkstemp(
                suffix=".html", prefix="sd_diff_report_", text=True
            )
            os.close(fd)

        with open(output_path, "w", encoding="utf-8") as f:
            f.write(html_content)

        return output_path

    @staticmethod
    def open_in_browser(file_path: str | Path) -> None:
        """Open an HTML file in the default browser."""
        webbrowser.open(f"file://{Path(file_path).resolve()}")

    # ------------------------------------------------------------------
    # .docx extraction (python-docx)
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_docx(file_path: str | Path, options: ExtractOptions) -> str:
        from docx import Document

        doc = Document(str(file_path))
        chunks: list[str] = []

        # Body paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                chunks.append(text)

        # Table cell paragraphs
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        text = para.text.strip()
                        if text:
                            chunks.append(text)

        # Headers (included only when ignore_header is False)
        if not options.ignore_header:
            for section in doc.sections:
                for para in section.header.paragraphs:
                    text = para.text.strip()
                    if text:
                        chunks.append(f"{t('doc.header_tag')} {text}")
                for para in section.first_page_header.paragraphs:
                    text = para.text.strip()
                    if text:
                        chunks.append(f"{t('doc.header_tag')} {text}")

        # Footers (included only when ignore_footer is False)
        if not options.ignore_footer:
            for section in doc.sections:
                for para in section.footer.paragraphs:
                    text = para.text.strip()
                    if text:
                        chunks.append(f"{t('doc.footer_tag')} {text}")
                for para in section.first_page_footer.paragraphs:
                    text = para.text.strip()
                    if text:
                        chunks.append(f"{t('doc.footer_tag')} {text}")

        return "\n".join(chunks)

    # ------------------------------------------------------------------
    # .doc extraction (pywin32 / win32com)
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_doc(file_path: str | Path, options: ExtractOptions) -> str:
        import win32com.client

        # Word header/footer story-type constants
        WD_HEADER_FOOTER_PRIMARY = 1
        WD_HEADER_FOOTER_FIRST_PAGE = 2
        WD_HEADER_FOOTER_EVEN_PAGES = 3

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False

        # Use absolute path — Word COM resolves relative paths against its
        # own default documents directory, which causes silent failures.
        abs_path = str(Path(file_path).resolve())
        doc = word.Documents.Open(abs_path, ReadOnly=True, Visible=False)
        try:
            chunks: list[str] = []

            # Main story (body) text
            story_text = doc.Content.Text.strip()
            if story_text:
                chunks.append(story_text)

            section_count = doc.Sections.Count

            # Headers (included only when ignore_header is False)
            if not options.ignore_header:
                for i in range(section_count):
                    section = doc.Sections(i + 1)
                    for hf_type in (
                        WD_HEADER_FOOTER_PRIMARY,
                        WD_HEADER_FOOTER_FIRST_PAGE,
                        WD_HEADER_FOOTER_EVEN_PAGES,
                    ):
                        try:
                            rng = section.Headers(hf_type).Range
                            text = rng.Text.strip()
                            if text:
                                chunks.append(f"{t('doc.header_tag')} {text}")
                        except Exception:  # noqa: BLE001, S110
                            pass

            # Footers (included only when ignore_footer is False)
            if not options.ignore_footer:
                for i in range(section_count):
                    section = doc.Sections(i + 1)
                    for hf_type in (
                        WD_HEADER_FOOTER_PRIMARY,
                        WD_HEADER_FOOTER_FIRST_PAGE,
                        WD_HEADER_FOOTER_EVEN_PAGES,
                    ):
                        try:
                            rng = section.Footers(hf_type).Range
                            text = rng.Text.strip()
                            if text:
                                chunks.append(f"{t('doc.footer_tag')} {text}")
                        except Exception:  # noqa: BLE001, S110
                            pass

            return "\n".join(chunks)
        finally:
            doc.Close(SaveChanges=False)
            word.Quit()

    # ------------------------------------------------------------------
    # .pdf extraction (pdfplumber)
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_pdf(file_path: str | Path, options: ExtractOptions) -> str:
        import pdfplumber

        chunks: list[str] = []

        with pdfplumber.open(str(file_path)) as pdf:
            for page in pdf.pages:
                page_width = page.width or 0
                page_height = page.height or 0

                top_crop = 0.0
                bottom_crop = 0.0

                # Crop header region (top 10% of page) if requested
                if options.ignore_header and page_height > 0:
                    top_crop = page_height * 0.10

                # Crop footer region (bottom 10% of page) if requested
                if options.ignore_footer and page_height > 0:
                    bottom_crop = page_height * 0.10

                if top_crop or bottom_crop:
                    crop_box = (
                        0,
                        top_crop,
                        page_width,
                        page_height - bottom_crop,
                    )
                    cropped = page.crop(crop_box)
                    text = cropped.extract_text()
                else:
                    text = page.extract_text()

                if text and text.strip():
                    chunks.append(text.strip())

        return "\n".join(chunks)


if __name__ == "__main__":
    print(t("doc.module_info"))
